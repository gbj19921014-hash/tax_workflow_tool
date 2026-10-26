import json
import ssl
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from desktop_app import TaxWorkflowApp, infer_tax_period
from exchange_rates import fetch_ecb_monthly_rates_to_eur
from run_workflow import italy_tax
from workflow_service import build_country_payload, write_outputs


class SeptemberFixesTest(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.config = json.loads(Path(__file__).with_name("config.json").read_text())

    def source(self, rows):
        path = self.root / "input.csv"
        pd.DataFrame(rows).to_csv(path, index=False, encoding="utf-8-sig")
        return path

    def row(self, **changes):
        row = {
            "ACTIVITY_PERIOD": "2026-SEP",
            "TAX_COLLECTION_RESPONSIBILITY": "SELLER",
            "TRANSACTION_TYPE": "SALE",
            "EXPORT_OUTSIDE_EU": "NO",
            "PRICE_OF_ITEMS_VAT_RATE_PERCENT": "0.22",
            "BUYER_VAT_NUMBER": "B67888438",
            "SALE_DEPART_COUNTRY": "IT",
            "SALE_ARRIVAL_COUNTRY": "ES",
            "TRANSACTION_CURRENCY_CODE": "EUR",
            "TOTAL_ACTIVITY_VALUE_AMT_VAT_INCL": "81.99",
            "TOTAL_ACTIVITY_VALUE_VAT_AMT": "14.78",
        }
        row.update(changes)
        return row

    def test_rate_input_accepts_unit_and_rejects_invalid_values(self):
        for text in ("0.2289", "0.2289 EUR", " 0.2289 eur "):
            self.assertEqual(TaxWorkflowApp._rate_value(None, text), 0.2289)
        self.assertIsNone(TaxWorkflowApp._rate_value(None, " "))
        for text in ("EUR", "NaN", "inf", "0", "-1", "0.2 GBP"):
            with self.subTest(text=text), self.assertRaises(ValueError):
                TaxWorkflowApp._rate_value(None, text)

    def test_ecb_request_keeps_certificate_verification(self):
        with patch("exchange_rates.urlopen") as request:
            request.return_value.__enter__.return_value.read.return_value = (
                b"TIME_PERIOD,OBS_VALUE\n2026-09,4.0\n"
            )
            result = fetch_ecb_monthly_rates_to_eur("2026-SEP", currencies=("PLN",))
        self.assertEqual(result["rates"]["PLN"], 0.25)
        context = request.call_args.kwargs["context"]
        self.assertEqual(context.verify_mode, ssl.CERT_REQUIRED)
        self.assertTrue(context.check_hostname)
        self.assertGreater(context.cert_store_stats()["x509_ca"], 0)

    def test_france_infers_unique_complete_quarter_with_extra_month(self):
        months = ["2026-JUN", "2026-JUL", "2026-AUG", "2026-SEP"]
        self.assertEqual(infer_tax_period("FR", months), "2026-Q3")
        self.assertIsNone(infer_tax_period("FR", ["2026-JUL", "2026-SEP"]))
        self.assertIsNone(infer_tax_period("FR", months + ["2026-APR", "2026-MAY"]))

    def test_france_excludes_other_months_even_when_they_match_tax_rules(self):
        rows = [self.row(ACTIVITY_PERIOD=month, SALE_DEPART_COUNTRY="FR",
                         SALE_ARRIVAL_COUNTRY="DE", BUYER_VAT_NUMBER="FR12345678901",
                         TOTAL_ACTIVITY_VALUE_AMT_VAT_INCL=amount)
                for month, amount in [("2026-JUL", "120"), ("2026-AUG", "60"),
                                      (" 2026-SEP ", "60"), ("2026-JUN", "1200")]]
        payload = build_country_payload(self.source(rows), "2026-Q3", "FR", self.config)
        self.assertEqual(payload["summary"][0]["命中行数"], 3)
        self.assertEqual(payload["final"][0]["金额EUR"], 240)
        self.assertEqual(payload["final"][1]["金额EUR"], 40)
        with self.assertRaisesRegex(ValueError, "不一致"):
            build_country_payload(self.source(rows[:2]), "2026-Q3", "FR", self.config)

    def test_italy_truncates_summary_and_final_tax_in_export(self):
        payload = build_country_payload(self.source([self.row()]), "2026-SEP", "IT", self.config)
        self.assertEqual(payload["summary"][2]["税金EUR"], 14.78)
        self.assertEqual(payload["final"][1]["金额EUR"], 14.78)
        self.assertEqual(italy_tax(-81.99, 0.22), -14.78)
        paths = write_outputs(payload, self.root / "output")
        workbook = load_workbook(paths["xlsx_path"], data_only=True)
        values = [cell.value for row in workbook["最终输出"] for cell in row]
        self.assertIn(14.78, values)
        self.assertNotIn(14.79, values)
        workbook.close()

    def test_poland_merges_overlap_without_removing_separate_source_rows(self):
        domestic = self.row(SALE_DEPART_COUNTRY="PL", SALE_ARRIVAL_COUNTRY="PL",
                            BUYER_VAT_NUMBER="PL5252733130", TOTAL_ACTIVITY_VALUE_AMT_VAT_INCL="123",
                            TOTAL_ACTIVITY_VALUE_VAT_AMT="23")
        cross_border = self.row(SALE_DEPART_COUNTRY="PL", SALE_ARRIVAL_COUNTRY="IT",
                                TOTAL_ACTIVITY_VALUE_AMT_VAT_INCL="246",
                                TOTAL_ACTIVITY_VALUE_VAT_AMT="46")
        payload = build_country_payload(self.source([domestic, domestic, cross_border]),
                                        "2026-SEP", "PL", self.config)
        self.assertEqual([row["命中行数"] for row in payload["summary"]], [0, 0, 2, 1])
        self.assertEqual(payload["summary"][3]["EUR金额总和"], 246)
        self.assertEqual([row["项目"] for row in payload["detail"]], ["2A", "2A", "3A"])
        self.assertEqual(payload["final"][0]["金额EUR"], 492)
        self.assertEqual(payload["final"][1]["金额EUR"], 92)
        paths = write_outputs(payload, self.root / "output")
        workbook = load_workbook(paths["xlsx_path"], data_only=True)
        self.assertEqual(workbook["命中明细"].max_row, 4)
        self.assertEqual(len(pd.read_csv(paths["output_dir"] / "命中明细.csv")), 3)
        workbook.close()

    def test_missing_hit_currency_rate_blocks_export(self):
        source = self.source([self.row(TRANSACTION_CURRENCY_CODE="PLN")])
        self.config["exchange_rates_to_eur"]["2026-SEP"] = {"PLN": None}
        payload = build_country_payload(source, "2026-SEP", "IT", self.config)
        output = self.root / "output"
        with self.assertRaisesRegex(ValueError, "PLN"):
            write_outputs(payload, output)
        self.assertFalse(output.exists())
        self.config["exchange_rates_to_eur"]["2026-SEP"] = {"PLN": 0.25}
        payload = build_country_payload(source, "2026-SEP", "IT", self.config)
        self.assertTrue(write_outputs(payload, output)["xlsx_path"].exists())

    def test_poland_complete_partial_and_no_overlap(self):
        domestic = self.row(SALE_DEPART_COUNTRY="PL", SALE_ARRIVAL_COUNTRY="PL",
                            BUYER_VAT_NUMBER="PL5252733130", TOTAL_ACTIVITY_VALUE_AMT_VAT_INCL="123",
                            TOTAL_ACTIVITY_VALUE_VAT_AMT="23")
        cross_border = dict(domestic, SALE_ARRIVAL_COUNTRY="IT",
                            TOTAL_ACTIVITY_VALUE_AMT_VAT_INCL="246")
        zero_vat_domestic = dict(domestic, TOTAL_ACTIVITY_VALUE_VAT_AMT="0")
        cases = [
            ("complete", [domestic], [1, 0], 123, 23),
            ("partial", [domestic, cross_border], [1, 1], 369, 69),
            ("none", [zero_vat_domestic, cross_border], [1, 1], 369, 69),
        ]
        for name, rows, counts, total, tax in cases:
            with self.subTest(overlap=name):
                payload = build_country_payload(self.source(rows), "2026-SEP", "PL", self.config)
                self.assertEqual([r["命中行数"] for r in payload["summary"]][2:], counts)
                self.assertEqual(len(payload["detail"]), len(rows))
                self.assertEqual(sum(r["行数"] for r in payload["currency"]), len(rows))
                self.assertEqual(payload["final"][0]["金额EUR"], total)
                self.assertEqual(payload["final"][1]["金额EUR"], tax)


if __name__ == "__main__":
    unittest.main()
